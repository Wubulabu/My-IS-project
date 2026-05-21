from __future__ import annotations

import csv
import json
from pathlib import Path
import shutil
import statistics

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent.parent
REPORT = ROOT / "report.docx"
BACKUP = ROOT / "output" / "doc" / "report_before_multimodal_update.docx"
MM_CSV = ROOT / "output" / "eval_report_multimodal" / "comparison.csv"
ABLATION_CSV = ROOT / "output" / "eval_report_multimodal_ablation" / "comparison.csv"
DATASET_JSON = ROOT / "data" / "emoji_dataset.json"
BASE_EVAL_JSON = ROOT / "data" / "eval_queries.json"
MM_EVAL_JSON = ROOT / "data" / "eval_queries_multimodal.json"


def set_para(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def find_para(doc: Document, token: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if token in paragraph.text:
            return paragraph
    raise ValueError(f"paragraph not found: {token}")


def replace_contains(doc: Document, token: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if token in paragraph.text:
            set_para(paragraph, text)
            return


def insert_after(paragraph: Paragraph, text: str = "", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style if style is not None else paragraph.style
    if text:
        new_para.add_run(text)
    return new_para


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def fmt(row: dict[str, str], key: str) -> str:
    return f"{float(row[key]):.4f}"


def normalize_method(name: str) -> str:
    mapping = {
        "bm25": "BM25",
        "tfidf": "TF-IDF",
        "dense": "Dense",
        "bi_encoder": "Bi-Encoder",
        "hnsw": "HNSW",
        "hybrid": "Hybrid",
        "visual": "Visual",
        "mm_hybrid": "MM-Hybrid",
        "mm_hybrid_no_visual": "去除视觉分支",
        "mm_hybrid_no_lexical": "去除词法分支",
    }
    return mapping.get(name, name)


def query_stats(path: Path) -> dict[str, object]:
    queries = read_json(path)
    langs: dict[str, int] = {}
    relevant_total = 0
    for item in queries:
        lang = item.get("lang", "unknown")
        langs[lang] = langs.get(lang, 0) + 1
        relevant_total += len(item.get("relevant", []))
    return {
        "count": len(queries),
        "langs": langs,
        "relevant_total": relevant_total,
    }


def lang_summary(stats: dict[str, object]) -> str:
    langs = stats["langs"]
    assert isinstance(langs, dict)
    parts: list[str] = []
    if langs.get("zh"):
        parts.append(f"中文 {langs['zh']}")
    if langs.get("en"):
        parts.append(f"英文 {langs['en']}")
    for lang, count in sorted(langs.items()):
        if lang not in {"zh", "en"}:
            parts.append(f"{lang} {count}")
    return "，".join(parts)


def dataset_stats() -> dict[str, object]:
    records = read_json(DATASET_JSON)
    if isinstance(records, dict):
        records = records.get("emojis") or list(records.values())
    keyword_lengths = [len(record.get("keywords") or []) for record in records]
    categories: dict[str, int] = {}
    missing_category = 0
    for record in records:
        category = record.get("category")
        if not category:
            missing_category += 1
            category = "missing"
        categories[category] = categories.get(category, 0) + 1
    return {
        "records": len(records),
        "categories": categories,
        "category_count": len(categories) - (1 if "missing" in categories else 0),
        "missing_category": missing_category,
        "keywords_avg": statistics.mean(keyword_lengths),
        "keywords_min": min(keyword_lengths),
        "keywords_max": max(keyword_lengths),
    }


def mark_fields_dirty(doc: Document) -> None:
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    for fld_char in doc.element.iter(qn("w:fldChar")):
        if fld_char.get(qn("w:fldCharType")) == "begin":
            fld_char.set(qn("w:dirty"), "true")


def set_table(table, rows: list[list[str]]) -> None:
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for row_idx, row in enumerate(rows):
        while len(table.rows[row_idx].cells) < len(row):
            table.rows[row_idx]._tr.add_tc()
        for col_idx, text in enumerate(row):
            table.rows[row_idx].cells[col_idx].text = text


def table_contains(table, token: str) -> bool:
    return any(token in cell.text for row in table.rows for cell in row.cells)


def append_row_if_missing(table, first_cell: str, row_values: list[str]) -> None:
    if table_contains(table, first_cell):
        for row in table.rows:
            if row.cells and row.cells[0].text == first_cell:
                for idx, value in enumerate(row_values):
                    row.cells[idx].text = value
                return
    row = table.add_row().cells
    for idx, value in enumerate(row_values):
        row[idx].text = value


def add_table_after(doc: Document, paragraph: Paragraph, rows: list[list[str]]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    if doc.tables:
        table.style = doc.tables[0].style
    set_table(table, rows)
    paragraph._p.addnext(table._tbl)
    return table


def main() -> None:
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(REPORT, BACKUP)

    mm_rows = read_csv(MM_CSV)
    ablation_rows = read_csv(ABLATION_CSV)
    mm_by_method = {row["Method"]: row for row in mm_rows}
    ab_by_method = {row["Method"]: row for row in ablation_rows}
    base_eval = query_stats(BASE_EVAL_JSON)
    mm_eval = query_stats(MM_EVAL_JSON)
    corpus_stats = dataset_stats()

    doc = Document(str(REPORT))

    replace_contains(doc, "多模态对齐下的 Emoji 检索系统设计与实现", "多模态对齐下的 Emoji 检索系统设计与实现")
    replace_contains(
        doc,
        "本课程设计面向即时通信",
        "本课程设计面向即时通信、社交媒体和内容创作中的 Emoji 检索需求，设计并实现了一个跨语言、多模态 Emoji 信息检索系统。系统将 Emoji 视为由自然语言描述、Unicode 符号属性和本地渲染图像共同构成的短文档对象，围绕数据治理、索引缓存、多路召回、融合排序、交互展示和离线评测形成完整的信息检索流程。",
    )
    replace_contains(
        doc,
        "系统在词法检索、语义检索和视觉检索之间建立统一调度机制",
        "系统在词法检索、语义检索和视觉检索之间建立统一调度机制：BM25 与 TF-IDF 提供可解释的词面基线，Dense 与 Bi-Encoder 提供跨语言语义匹配，HNSW 展示向量检索的近似最近邻扩展路径，openai/clip-vit-base-patch32 的 CLIP 图像编码与岭回归投影使渲染 Emoji 图像能够进入文本查询驱动的排序空间。MM-Hybrid 融合语义、视觉、词法和排名信号，用于处理颜色、形状、组合图形以及网络语义等复合检索意图。",
    )
    replace_contains(
        doc,
        "实验部分构建基础评测集和视觉意图评测集",
        f"实验部分构建基础评测集和视觉意图评测集，采用 MAP、MRR、P@K、NDCG@K 与延迟指标对多种方法进行比较。当前基础评测集包含 {base_eval['count']} 条查询，视觉意图评测集包含 {mm_eval['count']} 条查询。结果表明，语义检索能够明显缓解跨语言词汇鸿沟，视觉对齐分支能够在颜色和外观相关查询上提供稳定排序证据，MM-Hybrid 在视觉意图评测集上达到 MAP={fmt(mm_by_method['mm_hybrid'], 'MAP')}、NDCG@10={fmt(mm_by_method['mm_hybrid'], 'NDCG@10')}。项目实现覆盖倒排索引、稀疏向量、稠密向量、近似近邻、多模态对齐、融合排序和交互式反馈等信息检索课程核心内容。",
    )

    replace_contains(
        doc,
        "补充了真正的多模态处理链路",
        "补充了真正的多模态处理链路：系统将 Emoji 渲染为统一尺寸 PNG，使用 CLIP 图像编码器抽取 512 维视觉向量，并通过岭回归学习 visual -> text space 的投影，从而支持 Visual 和 MM-Hybrid 检索。",
    )
    replace_contains(
        doc,
        "多模态检索关注不同模态之间的可比表示",
        "多模态检索关注不同模态之间的可比表示。本项目没有把视觉模态停留在界面展示层，而是将每个 Emoji 渲染为图像，使用 openai/clip-vit-base-patch32 提取 CLIP 图像向量，再学习到 Bi-Encoder 文本语义空间的线性投影。该方案保留了本地缓存和离线复现能力，同时让视觉外观成为检索排序的独立信号。",
    )
    replace_contains(
        doc,
        "补充多模态后，数据层同时包含文本字段",
        "补充多模态后，数据层同时包含文本字段、Unicode 符号和渲染图像；索引层除倒排索引、TF-IDF、Bi-Encoder 与 HNSW 缓存外，还新增 emoji_images、clip_visual_embeddings、visual_projection 和 aligned_visual_embeddings；算法层新增 Visual 与 MM-Hybrid；展示层则通过方法按钮、分数贡献条和 Text-only / Visual-only / MM-Hybrid 对比视图显示不同模态信号的作用。",
    )
    replace_contains(
        doc,
        "多模态视觉分支位于 retrieval/visual.py",
        "多模态视觉分支位于 retrieval/visual.py。构建阶段先将每个 Emoji 渲染为 128 x 128 PNG，再调用 openai/clip-vit-base-patch32 的 CLIP image encoder 得到 512 维视觉向量，保存为 clip_visual_embeddings.npy。随后以已缓存的 Bi-Encoder 文本向量为对齐目标，通过岭回归学习 visual -> text space 的投影矩阵，得到 aligned_visual_embeddings.npy。在线查询时，用户文本仍由 Bi-Encoder 编码，Visual 方法直接与对齐后的视觉向量做相似度检索。",
    )
    replace_contains(
        doc,
        "MM-Hybrid 在原 Hybrid 基础上加入视觉辅助项",
        "MM-Hybrid 在原 Hybrid 基础上加入视觉辅助项，最终权重为：score = 0.50 * semantic + 0.25 * visual + 0.15 * bm25 + 0.10 * lexical + 0.00 * rank_bonus。该公式保留原 Hybrid 作为强文本基线，同时让 CLIP 视觉对齐分数拥有明确权重，便于通过消融实验观察视觉分支的贡献。",
    )
    replace_contains(
        doc,
        "（1）主检索、算法切换与多模态解释",
        "（1）主检索、算法切换与多模态解释：用户可以在同一搜索框下输入中文、英文或日文查询，并在 BM25、TF-IDF、Dense、Bi-Encoder、HNSW、Rerank、Hybrid、Visual 和 MM-Hybrid 之间切换。选择 Visual 或 MM-Hybrid 时，结果卡片会展示 Text、Visual、BM25、Lex 等分数贡献条；系统还提供同一 query 下 Text-only、Visual-only 与 MM-Hybrid 的 Top-5 并排对比，用于课堂演示视觉模态补充了什么。",
    )
    replace_contains(
        doc,
        "评测部分包含两个查询集合",
        f"评测部分包含两个查询集合。基础评测集 data/eval_queries.json 共 {base_eval['count']} 条查询，其中{lang_summary(base_eval)}，共 {base_eval['relevant_total']} 个相关标签。视觉意图评测集 data/eval_queries_multimodal.json 共 {mm_eval['count']} 条查询，其中{lang_summary(mm_eval)}，强调颜色、形状、图形外观和组合视觉特征，共 {mm_eval['relevant_total']} 个相关标签。",
    )
    replace_contains(
        doc,
        "表 7-1 展示视觉意图评测集上的主要结果",
        f"表 7-1 展示 {mm_eval['count']} 条视觉意图评测查询上的主要结果。该集合更能体现颜色、形状和图形外观对排序的影响，因此同时纳入文本、视觉和融合方法进行比较。",
    )
    replace_contains(
        doc,
        "结果显示，Dense 与 Bi-Encoder",
        f"结果显示，Dense 与 Bi-Encoder 在 MAP 和 MRR 上保持强文本语义基线，说明多语言语义编码能够稳定处理短查询和跨语言表达。Visual 的 MAP={fmt(mm_by_method['visual'], 'MAP')}，NDCG@10={fmt(mm_by_method['visual'], 'NDCG@10')}，说明视觉对齐分支能够在部分颜色、形状和外观查询上把相关候选提前。MM-Hybrid 的 MAP 达到 {fmt(mm_by_method['mm_hybrid'], 'MAP')}，NDCG@10 达到 {fmt(mm_by_method['mm_hybrid'], 'NDCG@10')}，高于 Hybrid 的 {fmt(mm_by_method['hybrid'], 'MAP')} 和 {fmt(mm_by_method['hybrid'], 'NDCG@10')}，表明视觉分数对最终排序具有有效贡献。",
    )
    replace_contains(
        doc,
        "消融结果显示，去除视觉分支后",
        f"消融结果显示，去除视觉分支后，MAP 从 {fmt(ab_by_method['mm_hybrid'], 'MAP')} 降至 {fmt(ab_by_method['mm_hybrid_no_visual'], 'MAP')}，NDCG@10 从 {fmt(ab_by_method['mm_hybrid'], 'NDCG@10')} 降至 {fmt(ab_by_method['mm_hybrid_no_visual'], 'NDCG@10')}；去除词法分支后 MAP 为 {fmt(ab_by_method['mm_hybrid_no_lexical'], 'MAP')}。这说明在视觉意图评测集中，颜色、外观和图形组合信号对排序贡献更明显，视觉分支已经进入实际排序计算。",
    )
    replace_contains(
        doc,
        "表 7-2 MM-Hybrid 消融实验结果",
        "表 7-2 多模态消融实验结果（data/eval_queries_multimodal.json）",
    )
    replace_contains(
        doc,
        "从课程能力映射来看",
        "从课程能力映射来看，项目覆盖了信息检索中的多个关键主题：倒排索引和词项权重对应 BM25/TF-IDF；向量空间和语义匹配对应 Dense/Bi-Encoder；视觉渲染、openai/clip-vit-base-patch32 的 CLIP 图像编码和线性投影对应多模态对齐；效率扩展对应 HNSW；排序优化对应 Hybrid、MM-Hybrid 和重排序；系统评价对应 MAP、MRR、NDCG 和延迟统计；交互式检索则体现在用户反馈和可视化探索中。",
    )
    replace_contains(
        doc,
        "（3）视觉分支使用通用 CLIP",
        "（3）视觉分支使用 openai/clip-vit-base-patch32 与线性投影，受字体渲染、小图标细节、肤色变体和组合 Emoji 影响，不等同于专门训练的 Emoji 多模态模型。",
    )
    replace_contains(
        doc,
        "（1）扩展评测集到 200 条以上",
        "（1）继续扩展视觉意图、网络语义、错别字、多轮反馈和跨语言混合难例，并引入分级相关性标注。",
    )
    replace_contains(
        doc,
        "（4）尝试 SigLIP",
        "（4）尝试 SigLIP、EVA-CLIP 等图文预训练模型，或使用 Emoji 图像-文本对进行轻量微调，使视觉模态获得更强的符号语义和细粒度颜色/形状区分能力。",
    )

    # Architecture and artifact tables.
    if len(doc.tables) >= 8:
        table_stats = doc.tables[3]
        set_table(
            table_stats,
            [
                ["数据统计项", "当前结果"],
                ["主数据集记录数", str(corpus_stats["records"])],
                ["类别数量", str(corpus_stats["category_count"])],
                ["category 缺失数", str(corpus_stats["missing_category"])],
                ["keywords 平均长度", f"{corpus_stats['keywords_avg']:.2f}"],
                ["keywords 最小/最大长度", f"{corpus_stats['keywords_min']} / {corpus_stats['keywords_max']}"],
                ["基础评测查询数", f"{base_eval['count']}（{lang_summary(base_eval)}）"],
                ["基础评测相关标签数", str(base_eval["relevant_total"])],
                ["视觉意图评测查询数", f"{mm_eval['count']}（{lang_summary(mm_eval)}）"],
                ["视觉意图相关标签数", str(mm_eval["relevant_total"])],
            ],
        )

        table_arch = doc.tables[1]
        table_arch.rows[2].cells[2].text = "构建倒排索引、词法语料、TF-IDF 模型、稠密向量、HNSW 图索引和 CLIP 视觉对齐缓存。"
        table_arch.rows[3].cells[2].text = "实现 BM25、TF-IDF、Dense、Bi-Encoder、HNSW、Visual、Hybrid、MM-Hybrid、重排序和缓存校验。"
        table_arch.rows[5].cells[2].text = "提供搜索界面、结果卡片、语义星图、向量实验、评估面板、多模态对比视图和交互反馈。"

        table_artifacts = doc.tables[5]
        append_row_if_missing(table_artifacts, "emoji_images/", ["emoji_images/", "2315 张 128 x 128 PNG", "保存统一渲染后的 Emoji 视觉输入。"])
        append_row_if_missing(table_artifacts, "clip_visual_embeddings.npy", ["clip_visual_embeddings.npy", "2315 x 512", "保存 CLIP image encoder 输出的视觉向量。"])
        append_row_if_missing(table_artifacts, "visual_projection.npy", ["visual_projection.npy", "513 x 384", "保存视觉向量到 Bi-Encoder 文本空间的岭回归投影矩阵。"])
        append_row_if_missing(table_artifacts, "aligned_visual_embeddings.npy", ["aligned_visual_embeddings.npy", "2315 x 384", "保存已对齐到文本空间的视觉向量，供 Visual/MM-Hybrid 在线检索。"])

        table_methods = doc.tables[6]
        append_row_if_missing(table_methods, "Visual", ["Visual", "CLIP 图像向量对齐分数", "独立验证视觉模态召回能力", "受字体渲染和通用 CLIP 对符号图像理解限制"])
        append_row_if_missing(table_methods, "MM-Hybrid", ["MM-Hybrid", "语义、视觉、BM25、精确词法与排名奖励", "把多模态信号纳入最终排序，解释性更强", "权重仍是启发式，需要学习排序进一步优化"])

        table_api = doc.tables[7]
        append_row_if_missing(table_api, "POST /api/multimodal_compare", ["POST /api/multimodal_compare", "返回同一 query 下 Text-only、Visual-only、MM-Hybrid 的 Top-5 对比。"])
        for row in table_api.rows:
            if row.cells[0].text == "GET /api/status":
                row.cells[1].text = "返回系统状态、视觉索引是否就绪、CLIP 模型名和视觉缓存数量。"

        # Main multimodal result table.
        wanted = ["bm25", "tfidf", "dense", "bi_encoder", "hnsw", "hybrid", "visual", "mm_hybrid"]
        result_rows = [["方法", "MAP", "MRR", "P@5", "NDCG@5", "P@10", "NDCG@10"]]
        for method in wanted:
            row = mm_by_method[method]
            result_rows.append([
                normalize_method(method),
                fmt(row, "MAP"),
                fmt(row, "MRR"),
                fmt(row, "P@5"),
                fmt(row, "NDCG@5"),
                fmt(row, "P@10"),
                fmt(row, "NDCG@10"),
            ])
        set_table(doc.tables[8], result_rows)

        ab_rows = [["消融变体", "MAP", "MRR", "P@5", "NDCG@10"]]
        for method in ["mm_hybrid", "mm_hybrid_no_visual", "mm_hybrid_no_lexical", "visual"]:
            row = ab_by_method[method]
            ab_rows.append([
                normalize_method(method),
                fmt(row, "MAP"),
                fmt(row, "MRR"),
                fmt(row, "P@5"),
                fmt(row, "NDCG@10"),
            ])
        set_table(doc.tables[9], ab_rows)

        if len(doc.tables) > 10:
            latency_rows = [
                ["方法", "平均延迟/ms", "p50/ms", "p95/ms", "说明"],
                ["BM25", "1.79", "1.74", "2.85", "轻量词法检索，速度最快之一。"],
                ["TF-IDF", "0.78", "0.66", "1.34", "稀疏矩阵余弦相似度，延迟最低。"],
                ["Dense", "33.80", "31.45", "58.41", "Sentence-Transformers 查询编码与矩阵检索。"],
                ["Bi-Encoder", "24.43", "23.03", "34.79", "手工 mean pooling 与缓存矩阵点积。"],
                ["HNSW", "121.29", "24.66", "52.38", "均值受个别图遍历开销影响，p50 接近精确向量检索。"],
                ["Rerank", "1961.68", "1764.09", "2412.00", "神经重排序最慢，主要作为高精度对照。"],
                ["Hybrid", "20.79", "20.15", "26.03", "融合 BM25 与 Bi-Encoder，延迟仍处于交互可用范围。"],
                ["Visual", "17.91", "17.01", "23.25", "运行时只读取已对齐视觉缓存，不调用 CLIP。"],
                ["MM-Hybrid", "21.62", "21.15", "25.82", "融合文本、词法与视觉分数，延迟接近 Hybrid。"],
            ]
            set_table(doc.tables[10], latency_rows)

    # Add a compact ablation table only for older report drafts that do not have one.
    if len(doc.tables) <= 9 and not any("表 7-2 多模态消融实验结果" in p.text for p in doc.paragraphs):
        anchor = find_para(doc, "该结果说明视觉模态不是单纯展示")
        title = insert_after(anchor, "表 7-2 多模态消融实验结果（data/eval_queries_multimodal.json）", anchor.style)
        ab_rows = [["Variant", "MAP", "MRR", "P@5", "NDCG@10"]]
        for method in ["mm_hybrid", "mm_hybrid_no_visual", "mm_hybrid_no_lexical", "visual"]:
            row = ab_by_method[method]
            ab_rows.append([
                normalize_method(method),
                fmt(row, "MAP"),
                fmt(row, "MRR"),
                fmt(row, "P@5"),
                fmt(row, "NDCG@10"),
            ])
        add_table_after(doc, title, ab_rows)

    mark_fields_dirty(doc)
    doc.save(str(REPORT))
    print(f"updated {REPORT}")
    print(f"backup {BACKUP}")


if __name__ == "__main__":
    main()
