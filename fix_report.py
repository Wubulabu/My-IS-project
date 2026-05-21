import sys
from docx import Document
from docx.enum.text import WD_BREAK
import re

def add_page_breaks_after_headings(doc_path):
    doc = Document(doc_path)
    new_doc = Document()
    for para in doc.paragraphs:
        # Copy paragraph
        new_para = new_doc.add_paragraph(para.text)
        # Preserve style
        if para.style:
            new_para.style = para.style
        # If heading style, add page break after it
        if para.style and para.style.name.startswith('Heading'):
            run = new_para.add_run()
            run.add_break(WD_BREAK.PAGE)
        else:
            # Simple cleanup: replace multiple spaces with single space
            cleaned = re.sub(r" {2,}", " ", para.text)
            if cleaned != para.text:
                new_para.text = cleaned
    # Save back to same file
    new_doc.save(doc_path)

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print('Usage: python fix_report.py <docx_path>')
        sys.exit(1)
    path = sys.argv[1]
    add_page_breaks_after_headings(path)
    print(f'Successfully updated {path}')
